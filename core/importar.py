"""
Importar datos reales del supermercado (2026-10-14).

El supermercado solo tiene dos cosas: la ficha de cada producto y sus
ventas diarias. Esta lectura acepta la plantilla de Fresh Analytics
(FAs_datos_demo_supermercado.xlsx) o lo que exporte su propio sistema, en
Excel (.xlsx) o CSV, siempre que tenga encabezados reconocibles (código,
fecha, cantidad...). Si vienen ventas por ticket, se suman por día.

- Productos: se crean o actualizan como origen REAL (el proceso nocturno
  no les inventa ventas, inventario ni mermas).
- Ventas: una Venta por producto y por día. Si ese día ya estaba cargado,
  se reemplaza (volver a importar el mismo archivo no duplica nada).
- Si el producto ya tiene lotes contados en el anaquel, las ventas nuevas
  los descuentan con FEFO (core/inventario.py).
"""

import csv
import io
import re
import unicodedata
from collections import defaultdict
from datetime import date, datetime, time, timedelta
from decimal import Decimal

from django.db import transaction
from django.utils import timezone

from .inventario import SimuladorInventario
from .models import Inventario, Producto, Venta

MAX_ERRORES = 25
# Código de la fila de ejemplo de la plantilla FAs_datos_demo_supermercado.xlsx
CODIGO_EJEMPLO = "7401234567890"

# Vida útil por categoría cuando el archivo no la trae (se avisa para que
# se confirme con el encargado del departamento).
VIDA_UTIL_POR_CATEGORIA = {"LACTEOS": 10, "CARNES": 4, "FRUTAS": 7, "VERDURAS": 6, "PANADERIA": 3}

SINONIMOS = {
    "codigo": ["codigodebarras", "codigobarras", "codigo", "upc", "ean", "sku", "barcode", "codigoproducto"],
    "nombre": ["nombredelproducto", "nombreproducto", "nombre", "descripcion", "producto", "articulo"],
    "categoria": ["categoria", "departamento", "familia", "seccion"],
    "unidad": ["sevendepor", "unidaddemedida", "unidadmedida", "medida", "unidad", "tipodeventa"],
    "vida_util": ["vidautildias", "vidautil", "diasdevida", "duraciondias"],
    "proveedor": ["proveedor", "marca"],
    "precio_compra": ["preciodecompraq", "preciodecompra", "preciocompra", "costo", "costounitario"],
    "precio_venta": ["preciodeventaq", "preciodeventa", "precioventa", "precio", "preciounitario", "pvp"],
    "fecha": ["fecha", "dia", "fechadeventa", "fechaventa"],
    "cantidad": ["cantidadvendida", "cantidad", "unidadesvendidas", "unidades", "libras", "librasvendidas", "peso"],
    "promocion": ["enpromocion", "promocion", "oferta"],
}

CATEGORIAS = {
    "lacteos": "LACTEOS", "lacteo": "LACTEOS", "carnes": "CARNES", "carne": "CARNES",
    "frutas": "FRUTAS", "fruta": "FRUTAS", "verduras": "VERDURAS", "verdura": "VERDURAS",
    "panaderia": "PANADERIA", "pan": "PANADERIA",
}


def _normalizar(texto):
    texto = unicodedata.normalize("NFKD", str(texto or "")).encode("ascii", "ignore").decode()
    return re.sub(r"[^a-z0-9]", "", texto.lower())


def _mapear_columnas(encabezados):
    """{campo: índice de columna} a partir de los encabezados del archivo."""
    normalizados = [_normalizar(h) for h in encabezados]
    columnas = {}
    for campo, opciones in SINONIMOS.items():
        for opcion in opciones:  # el sinónimo más específico gana
            if opcion in normalizados and normalizados.index(opcion) not in columnas.values():
                columnas[campo] = normalizados.index(opcion)
                break
    return columnas


def _texto(valor):
    if valor is None:
        return ""
    if isinstance(valor, float) and valor.is_integer():
        valor = int(valor)
    return str(valor).strip()


def _numero(valor):
    if valor in (None, ""):
        return None
    if isinstance(valor, (int, float, Decimal)):
        return float(valor)
    texto = str(valor).strip().replace("Q", "").replace(" ", "")
    if "," in texto and "." in texto:  # 1.234,56 o 1,234.56
        texto = texto.replace(".", "").replace(",", ".") if texto.rfind(",") > texto.rfind(".") else texto.replace(",", "")
    else:
        texto = texto.replace(",", ".")
    try:
        return float(texto)
    except ValueError:
        return None


def _fecha(valor):
    if isinstance(valor, datetime):
        return valor.date()
    if isinstance(valor, date):
        return valor
    texto = _texto(valor)
    for formato in ("%d/%m/%Y", "%d-%m-%Y", "%Y-%m-%d", "%d/%m/%y", "%Y/%m/%d"):
        try:
            return datetime.strptime(texto.split(" ")[0], formato).date()
        except ValueError:
            continue
    return None


def _si(valor):
    return _normalizar(valor) in ("si", "s", "yes", "1", "true", "x")


def leer_tablas(archivo, nombre):
    """Devuelve una lista de (nombre_de_hoja, filas) con la primera fila
    como encabezados."""
    if nombre.lower().endswith(".csv"):
        contenido = archivo.read()
        if isinstance(contenido, bytes):
            for codificacion in ("utf-8-sig", "latin-1"):
                try:
                    contenido = contenido.decode(codificacion)
                    break
                except UnicodeDecodeError:
                    continue
        dialecto = csv.Sniffer().sniff(contenido[:4096], delimiters=",;\t")
        return [("CSV", list(csv.reader(io.StringIO(contenido), dialecto)))]
    from openpyxl import load_workbook
    libro = load_workbook(archivo, read_only=True, data_only=True)
    return [(hoja.title, [list(f) for f in hoja.iter_rows(values_only=True)]) for hoja in libro.worksheets]


class Resultado:
    def __init__(self):
        self.errores, self.avisos = [], []
        self.productos_creados, self.productos_actualizados = [], []
        self.ventas = {}  # nombre -> {"dias", "desde", "hasta", "total", "unidad"}
        self.descontado_de_lotes = 0

    def error(self, texto):
        if len(self.errores) < MAX_ERRORES:
            self.errores.append(texto)

    @property
    def total_dias(self):
        return sum(v["dias"] for v in self.ventas.values())


def _es_ejemplo(fila):
    return any(
        isinstance(c, (str, int, float)) and ("ejemplo" in _normalizar(c) or _texto(c) == CODIGO_EJEMPLO)
        for c in fila
    )


def _importar_productos(hoja, filas, cols, res):
    obligatorios = ["codigo", "nombre", "categoria", "unidad", "precio_venta"]
    faltan = [c for c in obligatorios if c not in cols]
    if faltan:
        res.error(f"Hoja {hoja}: faltan columnas de producto ({', '.join(faltan)}).")
        return
    for n, fila in enumerate(filas[1:], start=2):
        if not any(c not in (None, "") for c in fila) or _es_ejemplo(fila):
            continue
        valor = lambda campo: fila[cols[campo]] if campo in cols and cols[campo] < len(fila) else None
        codigo, nombre = _texto(valor("codigo")), _texto(valor("nombre"))
        categoria = CATEGORIAS.get(_normalizar(valor("categoria")))
        unidad_txt = _normalizar(valor("unidad"))
        unidad = "PESO" if unidad_txt in ("libra", "libras", "lb", "peso", "pesovariable", "kg") else (
            "UNIDAD" if unidad_txt in ("unidad", "unidades", "pieza", "u", "c/u", "cu") else None)
        precio_venta = _numero(valor("precio_venta"))
        precio_compra = _numero(valor("precio_compra"))
        if not codigo or not nombre:
            res.error(f"Hoja {hoja}, fila {n}: falta el código o el nombre.")
            continue
        if not categoria:
            res.error(f"Hoja {hoja}, fila {n} ({nombre}): categoría \"{_texto(valor('categoria'))}\" no reconocida (Lácteos, Carnes, Frutas, Verduras o Panadería).")
            continue
        if not unidad:
            res.error(f"Hoja {hoja}, fila {n} ({nombre}): indica si se vende por Unidad o por Libra.")
            continue
        if precio_venta is None:
            res.error(f"Hoja {hoja}, fila {n} ({nombre}): falta el precio de venta.")
            continue
        vida = _numero(valor("vida_util"))
        if not vida:
            vida = VIDA_UTIL_POR_CATEGORIA[categoria]
            res.avisos.append(f"{nombre}: no traía vida útil; se usaron {vida} días (confírmalo con el encargado).")
        if precio_compra is None:
            precio_compra = round(precio_venta * 0.7, 2)
            res.avisos.append(f"{nombre}: no traía precio de compra; se estimó en 70% del de venta (Q {precio_compra:.2f}).")

        existente = Producto.objects.filter(codigo_upc=codigo).first()
        if existente and existente.origen != Producto.Origen.REAL:
            res.error(f"Hoja {hoja}, fila {n}: el código {codigo} ya es del producto de prueba \"{existente.nombre}\".")
            continue
        datos = dict(
            nombre=nombre[:200], categoria=categoria, unidad_medida=unidad, vida_util_dias=int(vida),
            proveedor=_texto(valor("proveedor"))[:100], precio_compra=Decimal(f"{precio_compra:.2f}"),
            precio_venta=Decimal(f"{precio_venta:.2f}"), activo=True, origen=Producto.Origen.REAL,
        )
        if existente:
            for campo, v in datos.items():
                setattr(existente, campo, v)
            existente.save()
            res.productos_actualizados.append(nombre)
        else:
            Producto.objects.create(codigo_upc=codigo, **datos)
            res.productos_creados.append(nombre)


def _importar_ventas(hoja, filas, cols, res):
    faltan = [c for c in ("codigo", "fecha", "cantidad") if c not in cols]
    if faltan:
        res.error(f"Hoja {hoja}: faltan columnas de venta ({', '.join(faltan)}).")
        return
    por_dia = defaultdict(lambda: {"cantidad": 0.0, "importe": 0.0, "con_precio": 0.0, "promo": False})
    hoy = timezone.localdate()
    productos = {}
    for n, fila in enumerate(filas[1:], start=2):
        if not any(c not in (None, "") for c in fila) or _es_ejemplo(fila):
            continue
        valor = lambda campo: fila[cols[campo]] if campo in cols and cols[campo] < len(fila) else None
        codigo = _texto(valor("codigo"))
        if not codigo and valor("fecha") in (None, ""):
            continue  # notas u otras filas que no son ventas
        if codigo not in productos:
            productos[codigo] = Producto.objects.filter(codigo_upc=codigo).first()
        producto = productos[codigo]
        if producto is None:
            res.error(f"Hoja {hoja}, fila {n}: el código {codigo or '(vacío)'} no está en la hoja de productos ni en FAs.")
            continue
        if producto.origen != Producto.Origen.REAL:
            res.error(f"Hoja {hoja}, fila {n}: {producto.nombre} es un producto de prueba; solo se importan ventas de productos reales.")
            continue
        fecha, cantidad = _fecha(valor("fecha")), _numero(valor("cantidad"))
        if fecha is None:
            res.error(f"Hoja {hoja}, fila {n}: fecha \"{_texto(valor('fecha'))}\" no válida (usa día/mes/año).")
            continue
        if fecha > hoy:
            res.error(f"Hoja {hoja}, fila {n}: la fecha {fecha:%d/%m/%Y} es futura.")
            continue
        if cantidad is None or cantidad < 0:
            res.error(f"Hoja {hoja}, fila {n}: cantidad no válida.")
            continue
        dia = por_dia[(producto.id_producto, fecha)]
        dia["cantidad"] += cantidad
        precio = _numero(valor("precio_venta"))
        if precio is not None:
            dia["importe"] += precio * cantidad
            dia["con_precio"] += cantidad
        dia["promo"] = dia["promo"] or _si(valor("promocion"))

    if not por_dia:
        return
    por_producto = defaultdict(dict)
    for (pid, fecha), dia in por_dia.items():
        por_producto[pid][fecha] = dia
    for pid, dias in por_producto.items():
        producto = Producto.objects.get(id_producto=pid)
        fechas = sorted(dias)
        ya_cargados = set(
            timezone.localtime(f).date() for f in Venta.objects.filter(
                producto=producto, fecha__date__gte=fechas[0] - timedelta(days=1),
                fecha__date__lte=fechas[-1] + timedelta(days=1),
            ).values_list("fecha", flat=True)
        ) & set(fechas)
        # Reemplaza los días que ya estaban (reimportar no duplica).
        for f in ya_cargados:
            Venta.objects.filter(producto=producto, fecha__date=f).delete()
        nuevas = []
        for f in fechas:
            dia = dias[f]
            cantidad = producto.redondear_cantidad(dia["cantidad"])
            precio = dia["importe"] / dia["con_precio"] if dia["con_precio"] else float(producto.precio_venta)
            nuevas.append(Venta(
                producto=producto, cantidad=cantidad, promocion_aplicada=dia["promo"],
                precio_unitario=Decimal(f"{precio:.2f}"),
                fecha=timezone.make_aware(datetime.combine(f, time(12))),
            ))
        Venta.objects.bulk_create(nuevas, batch_size=1000)

        # Ventas de días nuevos descuentan los lotes que ya estaban
        # contados ese día (FEFO). Los días históricos, anteriores al
        # conteo, no tienen lotes vigentes y no descuentan nada.
        nuevos = [f for f in fechas if f not in ya_cargados]
        if nuevos and Inventario.objects.filter(producto=producto, cantidad__gt=0).exists():
            sim = SimuladorInventario(Inventario)
            sim.cargar([pid], nuevos[0])
            for f in nuevos:
                res.descontado_de_lotes += sim.consumir(producto, dias[f]["cantidad"], f)
            sim.guardar()

        res.ventas[producto.nombre] = {
            "dias": len(fechas), "desde": fechas[0], "hasta": fechas[-1],
            "total": sum(producto.redondear_cantidad(dias[f]["cantidad"]) for f in fechas),
            "unidad": "lb" if producto.unidad_medida == Producto.UnidadMedida.PESO else "u.",
        }


def importar(archivo, nombre):
    """Lee el archivo y carga productos y ventas en una sola transacción:
    si algo revienta, no queda nada a medias."""
    res = Resultado()
    try:
        tablas = leer_tablas(archivo, nombre)
    except Exception:
        res.error("No se pudo leer el archivo. Usa la plantilla de FAs en Excel (.xlsx) o un CSV.")
        return res

    con_productos, con_ventas = [], []
    for hoja, filas in tablas:
        if not filas:
            continue
        cols = _mapear_columnas(filas[0])
        if "fecha" in cols and "cantidad" in cols:
            con_ventas.append((hoja, filas, cols))
        elif "nombre" in cols and "codigo" in cols:
            con_productos.append((hoja, filas, cols))
    if not con_productos and not con_ventas:
        res.error("No se encontró ninguna tabla con encabezados reconocibles (código, nombre, fecha, cantidad...).")
        return res

    with transaction.atomic():
        for hoja, filas, cols in con_productos:  # primero los productos
            _importar_productos(hoja, filas, cols, res)
        for hoja, filas, cols in con_ventas:
            _importar_ventas(hoja, filas, cols, res)
    return res
